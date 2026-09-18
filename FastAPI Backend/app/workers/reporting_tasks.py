import asyncio
import hashlib
import json
from datetime import datetime
from typing import Optional
from uuid import UUID

from celery import shared_task
from sqlalchemy import select

from app.core.celery.app import celery_app
from app.core.database import AsyncSessionLocal
from app.core.exceptions import ValidationException
from app.core.logging import get_logger
from app.core.storage.azure_blob import get_azure_blob_service
from app.modules.reporting.models import (
    ReportDefinition,
    ReportJob,
    ReportOutput,
    ReportStatus,
    ReportFormat,
)
from app.modules.reporting.repository import ReportingRepository
from app.modules.outbox.service import OutboxService
from app.modules.outbox.models import OutboxEventType

logger = get_logger(__name__)


async def _get_tenant_db(tenant_id: UUID):
    """Get a database session with tenant context set."""
    async with AsyncSessionLocal() as db:
        await db.execute(f"SET LOCAL app.current_tenant = '{tenant_id}'")
        yield db


@shared_task(bind=True, max_retries=3, default_retry_delay=60, autoretry_for=(Exception,))
def generate_report_task(self, job_id: str):
    """Generate a report asynchronously."""
    job_uuid = UUID(job_id)
    return asyncio.run(_generate_report_async(job_uuid))


async def _generate_report_async(job_id: UUID):
    async with AsyncSessionLocal() as db:
        repository = ReportingRepository(db, None)
        job = await repository.get_report_job(job_id)
        
        if not job:
            logger.error("Report job not found", job_id=str(job_id))
            return

        tenant_id = job.tenant_id
        await db.execute(f"SET LOCAL app.current_tenant = '{tenant_id}'")

        try:
            job.status = ReportStatus.PROCESSING
            job.started_at = datetime.utcnow()
            job.progress = 10
            job.current_step = "Loading report definition"
            await db.flush()

            definition = await repository.get_report_definition(job.report_definition_id)
            if not definition:
                raise ValidationException("Report definition not found")

            job.progress = 30
            job.current_step = "Executing query"
            await db.flush()

            result_data = await _execute_report_query(db, definition, job.parameters)

            job.progress = 60
            job.current_step = "Formatting output"
            await db.flush()

            output_bytes, checksum = await _format_report_output(
                result_data, definition.output_format, definition.output_config
            )

            job.progress = 80
            job.current_step = "Storing result"
            await db.flush()

            blob_service = get_azure_blob_service()
            storage_key = f"{tenant_id}/reports/{job_id}/{definition.name}.{definition.output_format.value}"
            
            await blob_service.upload_blob(
                storage_key=storage_key,
                data=output_bytes,
                content_type=_get_content_type(definition.output_format),
            )

            output = ReportOutput(
                job_id=job_id,
                storage_key=storage_key,
                format=definition.output_format,
                size_bytes=len(output_bytes),
                checksum=checksum,
                expires_at=datetime.utcnow() + timedelta(days=30),
                tenant_id=tenant_id,
            )
            db.add(output)
            await db.flush()

            job.status = ReportStatus.COMPLETED
            job.progress = 100
            job.current_step = "Completed"
            job.completed_at = datetime.utcnow()
            job.result_ref = storage_key
            job.result_format = definition.output_format
            job.result_size = len(output_bytes)
            job.result_checksum = checksum
            await db.flush()

            await _emit_report_completed_event(db, job, definition, output)

            logger.info("Report generated successfully", job_id=str(job_id), tenant_id=str(tenant_id))

        except Exception as e:
            logger.error("Report generation failed", job_id=str(job_id), error=str(e))
            job.status = ReportStatus.FAILED
            job.error_message = str(e)
            job.error_details = {"exception": type(e).__name__}
            job.completed_at = datetime.utcnow()
            await db.flush()

            await _emit_report_failed_event(db, job, str(e))
            
            if self.request.retries < self.max_retries:
                raise self.retry(exc=e)
            
            job.status = ReportStatus.FAILED
            job.error_details = {"exception": type(e).__name__, "max_retries_exceeded": True}
            await db.flush()

        await db.commit()


async def _execute_report_query(db, definition: ReportDefinition, parameters: dict) -> list[dict]:
    """Execute the report query based on the definition's query_config."""
    query_config = definition.query_config
    query_type = query_config.get("type", "sql")
    
    if query_type == "sql":
        sql = query_config.get("sql", "")
        if not sql:
            raise ValidationException("No SQL query defined in report definition")
        
        param_values = {**parameters}
        for key, value in param_values.items():
            if isinstance(value, str):
                param_values[key] = value.replace("'", "''")
        
        try:
            formatted_sql = sql.format(**param_values)
        except KeyError as e:
            raise ValidationException(f"Missing parameter: {e}")
        
        result = await db.execute(formatted_sql)
        columns = result.keys()
        return [dict(zip(columns, row)) for row in result.fetchall()]
    
    elif query_type == "orm":
        model_name = query_config.get("model")
        filters = query_config.get("filters", {})
        fields = query_config.get("fields", ["*"])
        
        if not model_name:
            raise ValidationException("No model specified in ORM query config")
        
        model_map = {
            "Client": "app.modules.clients.models.Client",
            "Matter": "app.modules.matters.models.Matter",
            "Task": "app.modules.tasks.models.Task",
            "Document": "app.modules.documents.models.Document",
            "Invoice": "app.modules.billing.models.Invoice",
            "Payment": "app.modules.billing.models.Payment",
            "Expense": "app.modules.billing.models.Expense",
            "ComplianceCycle": "app.modules.compliance.models.ComplianceCycle",
        }
        
        if model_name not in model_map:
            raise ValidationException(f"Unknown model: {model_name}")
        
        module_path, class_name = model_map[model_name].rsplit(".", 1)
        module = __import__(module_path, fromlist=[class_name])
        Model = getattr(module, class_name)
        
        query = select(Model).where(Model.tenant_id == definition.tenant_id)
        
        for field, value in filters.items():
            if hasattr(Model, field):
                query = query.where(getattr(Model, field) == value)
        
        if fields != ["*"]:
            query = query.with_entities(*[getattr(Model, f) for f in fields])
        
        result = await db.execute(query)
        if fields != ["*"]:
            return [dict(zip(fields, row)) for row in result.fetchall()]
        else:
            return [row._asdict() for row in result.fetchall()]
    
    else:
        raise ValidationException(f"Unknown query type: {query_type}")


async def _format_report_output(data: list[dict], format: ReportFormat, config: dict) -> tuple[bytes, str]:
    """Format the report data according to the specified format."""
    if format == ReportFormat.JSON:
        output = json.dumps(data, default=str, indent=2).encode("utf-8")
        checksum = hashlib.sha256(output).hexdigest()
        return output, checksum
    
    elif format == ReportFormat.CSV:
        import csv
        import io
        
        output_io = io.StringIO()
        if data:
            writer = csv.DictWriter(output_io, fieldnames=data[0].keys())
            writer.writeheader()
            writer.writerows(data)
        output = output_io.getvalue().encode("utf-8")
        checksum = hashlib.sha256(output).hexdigest()
        return output, checksum
    
    elif format == ReportFormat.EXCEL:
        import io
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        
        wb = Workbook()
        ws = wb.active
        ws.title = config.get("sheet_name", "Report")
        
        if data:
            headers = list(data[0].keys())
            header_font = Font(bold=True)
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font_color = Font(color="FFFFFF", bold=True)
            
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font_color
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
            
            for row_idx, row_data in enumerate(data, 2):
                for col_idx, header in enumerate(headers, 1):
                    value = row_data.get(header, "")
                    ws.cell(row=row_idx, column=col_idx, value=value)
            
            for column_cells in ws.columns:
                max_length = max(len(str(cell.value or "")) for cell in column_cells)
                ws.column_dimensions[column_cells[0].column_letter].width = min(max_length + 2, 50)
        
        output_io = io.BytesIO()
        wb.save(output_io)
        output = output_io.getvalue()
        checksum = hashlib.sha256(output).hexdigest()
        return output, checksum
    
    elif format == ReportFormat.PDF:
        from fpdf import FPDF
        import io
        
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, config.get("title", "Report"), ln=True, align="C")
        pdf.ln(5)
        
        if data:
            pdf.set_font("Helvetica", "B", 8)
            headers = list(data[0].keys())
            col_width = 190 / len(headers) if headers else 190
            
            for header in headers:
                pdf.cell(col_width, 8, str(header)[:20], border=1, align="C")
            pdf.ln()
            
            pdf.set_font("Helvetica", "", 7)
            for row_data in data[:100]:
                for header in headers:
                    value = str(row_data.get(header, ""))[:30]
                    pdf.cell(col_width, 6, value, border=1, align="L")
                pdf.ln()
        
        output = pdf.output(dest="S").encode("latin-1")
        checksum = hashlib.sha256(output).hexdigest()
        return output, checksum
    
    else:
        raise ValidationException(f"Unsupported output format: {format}")


def _get_content_type(format: ReportFormat) -> str:
    content_types = {
        ReportFormat.PDF: "application/pdf",
        ReportFormat.EXCEL: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ReportFormat.CSV: "text/csv",
        ReportFormat.JSON: "application/json",
    }
    return content_types.get(format, "application/octet-stream")


async def _emit_report_completed_event(db, job: ReportJob, definition: ReportDefinition, output: ReportOutput):
    """Emit outbox event for report completion."""
    outbox_service = OutboxService(db)
    await outbox_service.create_event(
        event_type=OutboxEventType.REPORT_COMPLETED,
        aggregate_type="report_job",
        aggregate_id=job.id,
        payload={
            "job_id": str(job.id),
            "definition_id": str(definition.id),
            "definition_name": definition.name,
            "output_id": str(output.id),
            "storage_key": output.storage_key,
            "format": output.format.value,
            "size_bytes": output.size_bytes,
            "checksum": output.checksum,
        },
        tenant_id=job.tenant_id,
    )


async def _emit_report_failed_event(db, job: ReportJob, error: str):
    """Emit outbox event for report failure."""
    outbox_service = OutboxService(db)
    await outbox_service.create_event(
        event_type=OutboxEventType.REPORT_FAILED,
        aggregate_type="report_job",
        aggregate_id=job.id,
        payload={
            "job_id": str(job.id),
            "definition_id": str(job.report_definition_id),
            "error": error,
        },
        tenant_id=job.tenant_id,
    )


@shared_task
def process_scheduled_reports():
    """Process due report schedules - called by Celery Beat."""
    return asyncio.run(_process_scheduled_reports_async())


async def _process_scheduled_reports_async():
    async with AsyncSessionLocal() as db:
        from app.modules.reporting.repository import ReportingRepository
        from app.modules.reporting.models import ReportSchedule
        
        repository = ReportingRepository(db, None)
        
        due_schedules = await _get_due_schedules(db)
        
        for schedule in due_schedules:
            tenant_id = schedule.tenant_id
            await db.execute(f"SET LOCAL app.current_tenant = '{tenant_id}'")
            
            try:
                job = ReportJob(
                    report_definition_id=schedule.report_definition_id,
                    parameters=schedule.parameters,
                    status=ReportStatus.QUEUED,
                    idempotency_key=f"schedule-{schedule.id}-{int(datetime.utcnow().timestamp())}",
                    tenant_id=tenant_id,
                )
                db.add(job)
                await db.flush()
                
                schedule.last_run_at = datetime.utcnow()
                schedule.last_run_status = "queued"
                
                from app.workers.reporting_tasks import generate_report_task
                generate_report_task.delay(str(job.id))
                
            except Exception as e:
                logger.error("Failed to queue scheduled report", schedule_id=str(schedule.id), error=str(e))
                schedule.last_run_status = f"failed: {str(e)[:100]}"
            
            await db.commit()


async def _get_due_schedules(db):
    """Get all due report schedules across all tenants."""
    from app.modules.firms.models import Firm
    
    result = await db.execute(select(Firm.id).where(Firm.is_active == True))
    tenant_ids = [row[0] for row in result.fetchall()]
    
    all_schedules = []
    for tenant_id in tenant_ids:
        await db.execute(f"SET LOCAL app.current_tenant = '{tenant_id}'")
        repository = ReportingRepository(db, tenant_id)
        schedules = await repository.get_due_schedules(datetime.utcnow())
        all_schedules.extend(schedules)
    
    return all_schedules


from datetime import timedelta